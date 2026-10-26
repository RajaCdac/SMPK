"""
Mark saved consolidations paid from an uploaded Excel.

A row is marked only when employee code, last pay, M1@277, M2@277,
and M1@359 / M2@359 (when those columns are in the sheet) match a
saved methodology2 or old-age consolidation. Pension figures are not
changed. Unmatched rows are written to a log file.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from M2_oldage_arrear.models import M2OldageArrearConsolidation
from methodology2.models import Methodology2Consolidation
from methodology2.services.bulk_excel_service import _clean_emp_cd
from methodology2.services.consolidation_report_service import _m2_m1_for_cpi

_LOG_HEADERS = (
    "Excel row",
    "Employee Code",
    "Name",
    "Source",
    "Field",
    "Excel amount",
    "Saved amount",
    "Reason",
)

_SOURCES = (
    ("oa", "Old age", M2OldageArrearConsolidation),
    ("met2", "Methodology 2", Methodology2Consolidation),
)


def _header_key(name) -> str:
    text = str(name or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _find_column(columns: list[str], predicate) -> str | None:
    for name in columns:
        if predicate(_header_key(name)):
            return name
    return None


def _is_m2(key: str) -> bool:
    return (
        "m ii" in key
        or "m2" in key
        or "methodology 2" in key
        or "methodology ii" in key
    )


def _is_m1(key: str) -> bool:
    if _is_m2(key):
        return False
    return (
        "m i" in key
        or "m1" in key
        or "methodology 1" in key
        or "methodology i" in key
    )


def _cpi_in(key: str, cpi: str) -> bool:
    return cpi in key.split()


def _locate_columns(columns: list[str]) -> dict:
    def last_pay(key: str) -> bool:
        if "scale" in key or "date" in key:
            return False
        if "pay revised" in key and "retirement" in key:
            return True
        return key in ("last pay", "pay revised", "revised pay")

    def phase(key: str) -> bool:
        return key == "phase" or key.startswith("phase ")

    def abstract(key: str) -> bool:
        return "abstract" in key

    def pay_date(key: str) -> bool:
        if key in ("date", "payment date", "abstract date", "paid date", "phase date"):
            return True
        if "date" not in key:
            return False
        blocked = ("retirement", "birth", "death", "dob", "doj", "from", "upto")
        return not any(word in key for word in blocked)

    found = {
        "emp": _find_column(
            columns,
            lambda key: key in ("employee code", "emp cd", "emp code", "empcd")
            or "employee code" in key,
        ),
        "name": _find_column(
            columns,
            lambda key: key in ("name", "pensioner name", "employee name"),
        ),
        "last_pay": _find_column(columns, last_pay),
        "m1_277": _find_column(
            columns, lambda key: _is_m1(key) and _cpi_in(key, "277")
        ),
        "m2_277": _find_column(
            columns, lambda key: _is_m2(key) and _cpi_in(key, "277")
        ),
        "m1_359": _find_column(
            columns, lambda key: _is_m1(key) and _cpi_in(key, "359")
        ),
        "m2_359": _find_column(
            columns, lambda key: _is_m2(key) and _cpi_in(key, "359")
        ),
        "phase": _find_column(columns, phase),
        "abstract": _find_column(columns, abstract),
        "pay_date": _find_column(columns, pay_date),
    }
    missing = [
        label
        for key, label in (
            ("emp", "Employee Code"),
            ("last_pay", "Pay (Revised) at the time of retirement"),
            ("m1_277", "Basic Pension amt as per M-I in 2017 (277 CPI)"),
            ("m2_277", "Revised Basic Pension amt as per M-II in 2017 (277 CPI)"),
        )
        if not found[key]
    ]
    if missing:
        raise ValueError(
            "This Excel is missing: " + ", ".join(missing) + "."
        )
    return found


def _blank(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return not text or text.lower() in ("nan", "none", "-")


def _amount(value):
    if _blank(value):
        return None
    text = str(value).strip().replace(",", "").replace("₹", "").replace("Rs.", "")
    text = text.replace("Rs", "").strip()
    try:
        return round(float(text), 2)
    except (TypeError, ValueError):
        return None


def _same(excel_amount, saved_amount) -> bool:
    if excel_amount is None and saved_amount is None:
        return True
    if excel_amount is None or saved_amount is None:
        return False
    return abs(float(excel_amount) - float(saved_amount)) < 0.005


def _show(amount) -> str:
    if amount is None:
        return ""
    number = float(amount)
    if abs(number - round(number)) < 0.001:
        return str(int(round(number)))
    return f"{number:.2f}"


def _text(value, limit: int = 40) -> str:
    if _blank(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text[:limit]


def _as_date(value):
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def _read_sheet(file_bytes: bytes) -> tuple[pd.DataFrame, int]:
    last_error = None
    for header in (0, 1):
        try:
            frame = pd.read_excel(BytesIO(file_bytes), header=header)
        except Exception as exc:
            last_error = exc
            continue
        columns = [str(c) for c in frame.columns]
        try:
            _locate_columns(columns)
        except ValueError:
            continue
        frame.columns = columns
        return frame, header
    if last_error:
        raise ValueError(f"Could not read the Excel file. {last_error}")
    raise ValueError(
        "Could not find Employee Code, last pay, M1@277, and M2@277 columns."
    )


def _saved_figures(obj) -> dict:
    m1_277, m2_277, _diff277 = _m2_m1_for_cpi(obj, 277)
    m1_359, m2_359, _diff359 = _m2_m1_for_cpi(obj, 359)
    return {
        "last_pay": _amount(getattr(obj, "last_pay", None)),
        "m1_277": _amount(m1_277),
        "m2_277": _amount(m2_277),
        "m1_359": _amount(m1_359),
        "m2_359": _amount(m2_359),
        "name": str(getattr(obj, "name", "") or getattr(obj, "pensioner_name", "") or ""),
    }


def _checks(excel: dict, saved: dict, cols: dict) -> list[tuple[str, object, object, bool]]:
    rows = [
        ("Last pay", excel["last_pay"], saved["last_pay"]),
        ("M1@277", excel["m1_277"], saved["m1_277"]),
        ("M2@277", excel["m2_277"], saved["m2_277"]),
    ]
    if cols.get("m1_359"):
        rows.append(("M1@359", excel["m1_359"], saved["m1_359"]))
    if cols.get("m2_359"):
        rows.append(("M2@359", excel["m2_359"], saved["m2_359"]))
    return [(label, excel_amt, saved_amt, _same(excel_amt, saved_amt)) for label, excel_amt, saved_amt in rows]


def _mark(obj, *, phase: str, abstract: str, pay_date, by_user: str) -> None:
    obj.payment_done = True
    obj.payment_done_at = timezone.now()
    obj.payment_done_by = str(by_user or "").strip()[:80]
    obj.payment_phase = phase
    obj.payment_abstract_no = abstract
    obj.payment_ref_date = pay_date
    obj.save(
        update_fields=[
            "payment_done",
            "payment_done_at",
            "payment_done_by",
            "payment_phase",
            "payment_abstract_no",
            "payment_ref_date",
            "updated_at",
        ]
    )


def _write_log(rows: list[dict]) -> tuple[str, str]:
    stamp = timezone.localtime().strftime("%Y-%m-%d_%H%M")
    parent = Path(getattr(settings, "BULK_OUTPUT_DIR", Path.home() / "Desktop"))
    parent.mkdir(parents=True, exist_ok=True)
    path = parent / f"paid-unmatched-{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Unmatched"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col, title in enumerate(_LOG_HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for index, row in enumerate(rows, start=2):
        values = [
            row.get("excel_row") or "",
            row.get("emp_cd") or "",
            row.get("name") or "",
            row.get("source") or "",
            row.get("field") or "",
            row.get("excel_amount") or "",
            row.get("saved_amount") or "",
            row.get("reason") or "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(index, col, value)
    widths = (12, 16, 28, 16, 14, 16, 16, 28)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    display = getattr(settings, "display_bulk_output_dir", None)
    shown = display(path) if callable(display) else str(path)
    return str(path), shown


def mark_paid_from_excel(file_bytes: bytes, *, by_user: str = "") -> dict:
    frame, header_row = _read_sheet(file_bytes)
    cols = _locate_columns([str(c) for c in frame.columns])
    log_rows: list[dict] = []
    marked = []
    seen_blank = 0

    for index, series in frame.iterrows():
        excel_row = int(index) + header_row + 2
        emp = _clean_emp_cd(series.get(cols["emp"]))
        excel_name = _text(series.get(cols["name"]), 200) if cols.get("name") else ""
        phase = _text(series.get(cols["phase"])) if cols.get("phase") else ""
        abstract = _text(series.get(cols["abstract"])) if cols.get("abstract") else ""
        pay_date = _as_date(series.get(cols["pay_date"])) if cols.get("pay_date") else None
        excel = {
            "last_pay": _amount(series.get(cols["last_pay"])),
            "m1_277": _amount(series.get(cols["m1_277"])),
            "m2_277": _amount(series.get(cols["m2_277"])),
            "m1_359": _amount(series.get(cols["m1_359"])) if cols.get("m1_359") else None,
            "m2_359": _amount(series.get(cols["m2_359"])) if cols.get("m2_359") else None,
        }
        if not emp and all(value is None for value in excel.values()) and not phase:
            seen_blank += 1
            continue
        if not emp:
            log_rows.append(
                {
                    "excel_row": excel_row,
                    "emp_cd": "",
                    "name": excel_name,
                    "source": "",
                    "field": "Employee Code",
                    "excel_amount": "",
                    "saved_amount": "",
                    "reason": "Employee code is blank",
                }
            )
            continue

        failures = []
        matches = []
        for _code, source_label, model in _SOURCES:
            obj = model.objects.filter(emp_cd=emp).first()
            if not obj:
                continue
            saved = _saved_figures(obj)
            checks = _checks(excel, saved, cols)
            failed = [item for item in checks if not item[3]]
            if failed:
                for label, excel_amt, saved_amt, _ok in failed:
                    failures.append(
                        {
                            "excel_row": excel_row,
                            "emp_cd": emp,
                            "name": excel_name or saved["name"],
                            "source": source_label,
                            "field": label,
                            "excel_amount": _show(excel_amt),
                            "saved_amount": _show(saved_amt),
                            "reason": "Does not match",
                        }
                    )
                continue
            matches.append((obj, excel_name or saved["name"], source_label))

        if matches:
            for obj, name, source_label in matches:
                _mark(
                    obj,
                    phase=phase,
                    abstract=abstract,
                    pay_date=pay_date,
                    by_user=by_user,
                )
                marked.append(
                    {
                        "emp_cd": emp,
                        "name": name,
                        "source": source_label,
                        "phase": phase,
                        "abstract_no": abstract,
                        "payment_date": pay_date.isoformat() if pay_date else "",
                    }
                )
            continue

        if failures:
            log_rows.extend(failures)
        else:
            log_rows.append(
                {
                    "excel_row": excel_row,
                    "emp_cd": emp,
                    "name": excel_name,
                    "source": "",
                    "field": "",
                    "excel_amount": "",
                    "saved_amount": "",
                    "reason": "No saved consolidation",
                }
            )

    log_path = ""
    log_display = ""
    if log_rows:
        log_path, log_display = _write_log(log_rows)

    return {
        "marked_count": len(marked),
        "unmatched_count": len({(r["excel_row"], r["emp_cd"]) for r in log_rows}),
        "log_row_count": len(log_rows),
        "marked": marked,
        "unmatched": log_rows,
        "log_path": log_display or log_path,
        "compared_359": bool(cols.get("m1_359") or cols.get("m2_359")),
        "blank_rows_skipped": seen_blank,
    }
