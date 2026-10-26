"""
Pay & Cash arrear sheet for a list of saved consolidations.

Does not recalculate Methodology-2 and does not mark a row as paid.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from methodology2.services.arrear_payment_service import (
    calculate_consolidation_arrear,
    validate_arrear_period,
)

MAX_EMPLOYEES = 500

_FULL_MONTH = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)
_CLASS_ROMAN = {"1": "I", "2": "II", "3": "III", "4": "IV"}

_FIXED_HEADERS = (
    "Sl No",
    "Name",
    "Employee Code",
    "Retirement Date",
    "Class",
    "Designation",
    "Pay (Revised) at the time of retirement",
    "Pay Scale at the time of retirement",
    "TQS",
    "Basic Pension as per M-I",
    "Basic Pension as per M-II",
    "CPI",
    "Beneficial under",
    "Amt of Benefit (Excluding DA)",
)
_TAIL_HEADERS = (
    "Total bill amount",
    "Phase",
    "Arrear to be paid upto",
    "Abstract No.",
)

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=9)
_EMP_FILL = PatternFill("solid", fgColor="FFF2CC")
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")
_LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")
_RIGHT = Alignment(vertical="center", horizontal="right")


def paid_upto_label(from_month, from_year) -> str:
    """Month before this payment starts. Paying from May 2026 is paid upto April 2026."""
    month = int(from_month)
    year = int(from_year)
    if month <= 1:
        month = 12
        year -= 1
    else:
        month -= 1
    return f"{_FULL_MONTH[month - 1]}, {year}"


def _pension_obj(emp: str):
    from M2_oldage_arrear.models import M2OldageArrearConsolidation
    from methodology2.models import Methodology2Consolidation

    met2 = Methodology2Consolidation.objects.filter(emp_cd=emp).first()
    if met2 is not None:
        return met2
    return M2OldageArrearConsolidation.objects.filter(emp_cd=emp).first()


def _num(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _dmy(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    text = str(value).strip()[:10]
    parts = text.split("-")
    if len(parts) == 3 and len(parts[0]) == 4:
        return f"{parts[2]}-{parts[1]}-{parts[0]}"
    return text


def _class_label(category) -> str:
    text = str(category or "").strip()
    return _CLASS_ROMAN.get(text, text)


def _scale_text(scale, retirement_date) -> str:
    from methodology2.services.class1_2_service import pay_scale_for_report

    return pay_scale_for_report(scale, retirement_date)


def _identity(emp: str) -> dict:
    obj = _pension_obj(emp)
    if obj is None:
        return {
            "name": "",
            "retirement_date": "",
            "class_label": "",
            "designation": "",
            "last_pay": None,
            "scale": "",
            "tqs": "",
        }
    retired = getattr(obj, "retirement_date", None)
    return {
        "name": str(getattr(obj, "name", "") or ""),
        "retirement_date": _dmy(retired),
        "class_label": _class_label(getattr(obj, "category", "")),
        "designation": str(getattr(obj, "designation", "") or ""),
        "last_pay": _num(getattr(obj, "last_pay", None)),
        "scale": _scale_text(getattr(obj, "scale", ""), retired),
        "tqs": str(getattr(obj, "tqs_display", "") or ""),
    }


def _beneficial(diff) -> str:
    if diff is None:
        return ""
    if diff > 0.005:
        return "M-II"
    if diff < -0.005:
        return "M-I"
    return "Same"


def _pair(lines: list[dict]) -> dict:
    """One M-I / M-II pair when every month in the window uses it."""
    if not lines:
        return {
            "split": False,
            "m1": None,
            "m2": None,
            "cpi": "",
            "benefit": None,
            "beneficial": "",
        }
    same = (
        len({line.get("m1") for line in lines}) == 1
        and len({line.get("m2") for line in lines}) == 1
        and len({line.get("pension_diff") for line in lines}) == 1
        and len({line.get("cpi") for line in lines}) == 1
    )
    if not same:
        return {
            "split": True,
            "m1": None,
            "m2": None,
            "cpi": "",
            "benefit": None,
            "beneficial": "Split",
        }
    diff = lines[0].get("pension_diff")
    return {
        "split": False,
        "m1": lines[0].get("m1"),
        "m2": lines[0].get("m2"),
        "cpi": lines[0].get("cpi") or "",
        "benefit": diff,
        "beneficial": _beneficial(diff),
    }


def _add_month(d: date) -> date:
    if d.month == 12:
        return date(d.year + 1, 1, 1)
    return date(d.year, d.month + 1, 1)


def _prev_month(d: date) -> date:
    if d.month == 1:
        return date(d.year - 1, 12, 1)
    return date(d.year, d.month - 1, 1)


def _month_end(d: date) -> date:
    from calendar import monthrange

    return date(d.year, d.month, monthrange(d.year, d.month)[1])


def _short_month(d: date) -> str:
    return f"{_FULL_MONTH[d.month - 1][:3]} {d.year}"


def _payment_periods(from_month, from_year, to_month, to_year) -> list[dict]:
    """
    One block per DA change inside the payment window.

    May–Jun stays one block and Jul–Sep stays the next when DA changes in July.
    An old-age start in June does not open another block.
    """
    from M2_oldage_arrear.services.da_service import da_wef_dates

    start = date(int(from_year), int(from_month), 1)
    end = date(int(to_year), int(to_month), 1)
    last = _month_end(end)
    points = {start}
    for grp in (1, 3):
        for wef in da_wef_dates(grp, start, last):
            points.add(date(wef.year, wef.month, 1))
    ordered = [p for p in sorted(points) if start <= p <= end]
    periods = []
    for index, begin in enumerate(ordered):
        finish = _prev_month(ordered[index + 1]) if index + 1 < len(ordered) else end
        if finish < begin:
            continue
        title = _short_month(begin) if begin == finish else f"{_short_month(begin)} - {_short_month(finish)}"
        periods.append(
            {
                "id": begin.isoformat(),
                "from": begin,
                "to": finish,
                "title": title,
            }
        )
    return periods


def _line_months(line: dict) -> list[date]:
    start = date.fromisoformat(str(line.get("from_iso"))[:10])
    end = date.fromisoformat(str(line.get("to_iso"))[:10])
    months = []
    cursor = start
    while cursor <= end:
        months.append(cursor)
        cursor = _add_month(cursor)
    return months


def _period_id(periods: list[dict], month: date) -> str | None:
    for period in periods:
        if period["from"] <= month <= period["to"]:
            return period["id"]
    return None


def _amounts_by_period(lines: list[dict], periods: list[dict]) -> dict:
    from methodology2.services.arrear_payment_service import _money

    totals = {period["id"]: {"oa": 0.0, "basic": 0.0, "da": 0.0} for period in periods}
    for line in lines or []:
        months = _line_months(line)
        if not months:
            continue
        ids = [_period_id(periods, month) for month in months]
        if ids and len(set(ids)) == 1 and ids[0]:
            pid = ids[0]
            count = int(line.get("months") or len(months))
            totals[pid]["oa"] = _money(
                totals[pid]["oa"] + float(line.get("oldage_benefit") or 0) * count
            )
            totals[pid]["basic"] = _money(
                totals[pid]["basic"] + float(line.get("pension_arrear") or 0)
            )
            totals[pid]["da"] = _money(
                totals[pid]["da"] + float(line.get("da_relief") or 0)
            )
    return totals


def _pct_text(value) -> str:
    number = float(value or 0)
    if number == int(number):
        return str(int(number))
    return f"{number:g}"


def _column_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _comparison_columns(emp: str) -> list[dict]:
    obj = _pension_obj(emp)
    raw = getattr(obj, "pension_comparison", None) if obj is not None else None
    if not isinstance(raw, dict):
        return []
    return [col for col in (raw.get("columns") or []) if isinstance(col, dict)]


def _column_covers_month(col: dict, month: date) -> bool:
    start = _column_date(col.get("period_from"))
    end = _column_date(col.get("period_to"))
    if start is not None and month < start:
        return False
    if end is not None and month > end:
        return False
    return True


def _slab_start(columns: list[dict], month: date, percent: float, cpi, kind):
    """When this old-age percent began, even if that month is before the payment."""
    matches = []
    for col in columns:
        try:
            col_pct = float(col.get("oldage_percent") or 0)
        except (TypeError, ValueError):
            continue
        if col_pct != percent or not _column_covers_month(col, month):
            continue
        matches.append(col)
    if not matches:
        return None

    def rank(col):
        return (
            str(col.get("cpi") or "") == str(cpi or ""),
            str(col.get("kind") or "") == str(kind or ""),
        )

    chosen = sorted(matches, key=rank)[-1]
    return _column_date(chosen.get("period_from"))


def _oab_from_note(emp: str, lines: list[dict]) -> str:
    """
    Start month and percent of each old-age rate used in the payment.

    The start can be before the payment. Emp 00130 is 30% through May–Sep 2026
    because that rate began in Nov 2025, so the cell is Nov 2025 @ 30%.
    """
    columns = _comparison_columns(emp)
    notes = []
    previous = None
    for line in lines or []:
        percent = float(line.get("oldage_percent") or 0)
        if percent == previous:
            continue
        previous = percent
        if percent <= 0:
            continue
        month = date.fromisoformat(str(line.get("from_iso"))[:10])
        started = _slab_start(
            columns, month, percent, line.get("cpi"), line.get("pension_kind")
        )
        label = _short_month(started) if started else (line.get("from_label") or "")
        notes.append(f"{label} @ {_pct_text(percent)}%")
    return ", ".join(notes)


def calculate_arrear_batch(emp_codes, from_month, from_year, to_month, to_year) -> dict:
    period_error = validate_arrear_period(from_month, from_year, to_month, to_year)
    if period_error:
        return {"error": period_error}

    codes: list[str] = []
    seen: set[str] = set()
    for raw in emp_codes or []:
        emp = str(raw or "").strip()[:5]
        if not emp or emp in seen:
            continue
        seen.add(emp)
        codes.append(emp)
    if not codes:
        return {"error": "Excel must include a column named EMP_CD."}
    if len(codes) > MAX_EMPLOYEES:
        return {"error": f"Upload at most {MAX_EMPLOYEES} employee codes at a time."}

    rows = []
    for emp in codes:
        calc = calculate_consolidation_arrear(emp, from_month, from_year, to_month, to_year)
        ident = _identity(emp)
        if calc.get("error"):
            rows.append({
                "ok": False,
                "emp_cd": emp,
                "error": calc["error"],
                **ident,
            })
            continue
        lines = calc.get("lines") or []
        rows.append({
            "ok": True,
            "error": "",
            **ident,
            "name": ident["name"] or calc.get("name") or "",
            "emp_cd": calc.get("emp_cd") or emp,
            "category": calc.get("category") or "",
            "class12": bool(calc.get("class12")),
            "pair": _pair(lines),
            "lines": lines,
            "pension_arrear": calc.get("pension_arrear"),
            "da_relief": calc.get("da_relief"),
            "total": calc.get("total"),
            "from_label": calc.get("from_label"),
            "to_label": calc.get("to_label"),
        })

    first_ok = next((row for row in rows if row.get("ok")), None)
    return {
        "from_label": first_ok.get("from_label") if first_ok else "",
        "to_label": first_ok.get("to_label") if first_ok else "",
        "paid_upto": paid_upto_label(from_month, from_year),
        "requested": len(codes),
        "ok_count": sum(1 for row in rows if row.get("ok")),
        "failed_count": sum(1 for row in rows if not row.get("ok")),
        "rows": rows,
    }


def _write_header(ws, periods: list[dict]) -> int:
    """Two header rows. Returns the first data column of the tail."""
    fixed = len(_FIXED_HEADERS)
    for col, title in enumerate(_FIXED_HEADERS, start=1):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        cell = ws.cell(1, col, title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cell.border = _THIN
        ws.cell(2, col).fill = _HEADER_FILL
        ws.cell(2, col).border = _THIN
        ws.cell(2, col).font = _HEADER_FONT

    col = fixed + 1
    ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    oab_head = ws.cell(1, col, "OAB from")
    oab_head.fill = _HEADER_FILL
    oab_head.font = _HEADER_FONT
    oab_head.alignment = _WRAP
    oab_head.border = _THIN
    ws.cell(2, col).fill = _HEADER_FILL
    ws.cell(2, col).border = _THIN
    ws.cell(2, col).font = _HEADER_FONT
    col += 1

    sub = ("Old age diff", "Basic + old age", "DA relief")
    for period in periods:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        head = ws.cell(1, col, period["title"])
        head.fill = _HEADER_FILL
        head.font = _HEADER_FONT
        head.alignment = _WRAP
        head.border = _THIN
        for offset, title in enumerate(sub):
            cell = ws.cell(2, col + offset, title)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _WRAP
            cell.border = _THIN
            ws.cell(1, col + offset).border = _THIN
            ws.cell(1, col + offset).fill = _HEADER_FILL
        col += 3

    tail_at = col
    for title in _TAIL_HEADERS:
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        cell = ws.cell(1, col, title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cell.border = _THIN
        ws.cell(2, col).fill = _HEADER_FILL
        ws.cell(2, col).border = _THIN
        ws.cell(2, col).font = _HEADER_FONT
        col += 1
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"
    return tail_at


def _put_money(cell, value):
    if value is None or value == "":
        cell.value = None
        cell.alignment = _RIGHT
        return
    cell.value = float(value)
    cell.number_format = "#,##0.00"
    cell.alignment = _RIGHT


def _write_row(ws, excel_row: int, sl_no: int, row: dict, periods: list[dict], phase: str, paid_upto: str):
    pair = row.get("pair") or {}
    values = [
        sl_no,
        row.get("name") or "",
        row.get("emp_cd") or "",
        row.get("retirement_date") or "",
        row.get("class_label") or "",
        row.get("designation") or "",
    ]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(excel_row, col, value)
        cell.alignment = _LEFT if col in (2, 6) else Alignment(vertical="center", horizontal="center")
        cell.border = _THIN
        if col == 3:
            cell.fill = _EMP_FILL
    _put_money(ws.cell(excel_row, 7), row.get("last_pay"))
    ws.cell(excel_row, 7).border = _THIN
    scale_cell = ws.cell(excel_row, 8, row.get("scale") or "")
    scale_cell.alignment = _LEFT
    scale_cell.border = _THIN
    tqs_cell = ws.cell(excel_row, 9, row.get("tqs") or "")
    tqs_cell.alignment = Alignment(vertical="center", horizontal="center")
    tqs_cell.border = _THIN
    _put_money(ws.cell(excel_row, 10), pair.get("m1"))
    _put_money(ws.cell(excel_row, 11), pair.get("m2"))
    ws.cell(excel_row, 10).border = _THIN
    ws.cell(excel_row, 11).border = _THIN
    cpi_cell = ws.cell(excel_row, 12, pair.get("cpi") or "")
    ben_cell = ws.cell(excel_row, 13, pair.get("beneficial") or "")
    cpi_cell.alignment = Alignment(vertical="center", horizontal="center")
    ben_cell.alignment = Alignment(vertical="center", horizontal="center")
    cpi_cell.border = _THIN
    ben_cell.border = _THIN
    _put_money(ws.cell(excel_row, 14), pair.get("benefit"))
    ws.cell(excel_row, 14).border = _THIN

    oab_cell = ws.cell(
        excel_row,
        len(_FIXED_HEADERS) + 1,
        _oab_from_note(row.get("emp_cd") or "", row.get("lines") or []),
    )
    oab_cell.alignment = Alignment(vertical="center", horizontal="center")
    oab_cell.border = _THIN

    amounts = _amounts_by_period(row.get("lines") or [], periods)
    col = len(_FIXED_HEADERS) + 2
    for period in periods:
        bucket = amounts.get(period["id"]) or {"oa": 0, "basic": 0, "da": 0}
        _put_money(ws.cell(excel_row, col), bucket["oa"])
        _put_money(ws.cell(excel_row, col + 1), bucket["basic"])
        _put_money(ws.cell(excel_row, col + 2), bucket["da"])
        for offset in range(3):
            ws.cell(excel_row, col + offset).border = _THIN
        col += 3

    _put_money(ws.cell(excel_row, col), row.get("total"))
    ws.cell(excel_row, col).border = _THIN
    phase_cell = ws.cell(excel_row, col + 1, phase or "")
    upto_cell = ws.cell(excel_row, col + 2, paid_upto)
    abstract = ws.cell(excel_row, col + 3, "")
    for cell in (phase_cell, upto_cell, abstract):
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = _THIN


def build_arrear_payment_excel_bytes(
    emp_codes,
    from_month,
    from_year,
    to_month,
    to_year,
    phase: str = "",
) -> tuple[bytes, str]:
    batch = calculate_arrear_batch(emp_codes, from_month, from_year, to_month, to_year)
    if batch.get("error"):
        raise ValueError(batch["error"])
    failed = [row for row in batch["rows"] if not row.get("ok")]
    if failed:
        first = failed[0]
        raise ValueError(f"{first.get('emp_cd')}: {first.get('error')}")
    ready = [row for row in batch["rows"] if row.get("ok")]
    if not ready:
        raise ValueError("None of the selected employees have an arrear calculation.")

    periods = _payment_periods(from_month, from_year, to_month, to_year)
    paid = batch["paid_upto"]
    phase_text = str(phase or "").strip()

    wb = Workbook()
    ws = wb.active
    ws.title = "Arrear"
    _write_header(ws, periods)
    for index, row in enumerate(ready, start=1):
        _write_row(ws, index + 2, index, row, periods, phase_text, paid)

    widths = {
        1: 8, 2: 28, 3: 14, 4: 16, 5: 10, 6: 22,
        7: 18, 8: 28, 9: 16, 10: 16, 11: 16, 12: 10, 13: 16, 14: 16,
    }
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = widths.get(col, 18)

    buf = BytesIO()
    wb.save(buf)
    start = date(int(from_year), int(from_month), 1)
    end = date(int(to_year), int(to_month), 1)
    filename = (
        f"Arrear_{start.strftime('%b%Y')}_{end.strftime('%b%Y')}.xlsx"
    )
    return buf.getvalue(), filename
