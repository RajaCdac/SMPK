"""Excel report from consolidation snapshots (M-I vs M-II, optional Old Age)."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from M2_oldage_arrear.services.oldage_print_columns import (
    oldage_slab_excel_headers,
    oldage_slab_excel_values,
)
from methodology2.models import Methodology2Consolidation, M2ConsolCase, M2ConsolRate

HEADERS_BASE = [
    "Sl. NO.",
    "Pensioner Name",
    "Case NO.",
    "Roll No.",
    "Employee Code",
    "Retirement Date",
    "Class",
    "Designation",
    "Pay (Revised) at the time of retirement",
    "Pay scale at the time of Retirement",
    "Date of Death",
    "Basic Pension amt as per M-I in 2017 (277 CPI)",
    "Revised Basic Pension amt as per M-II in 2017 (277 CPI)",
    "Beneficial under M-I or M-II? (277)",
    "Amt of Benefit (PM) excluding D.A (277)",
    "Basic Pension amt as per M-I in 2022 (359 CPI)",
    "Revised Basic Pension amt as per M-II in 2022 (359 CPI)",
    "Beneficial under M-I or M-II? (359)",
    "Amt of Benefit (PM) excluding D.A (359)",
]

HEADERS_ENHANCED = [
    "Enhanced / Double FP upto",
    "Family pension from",
    "Has death / enhanced split?",
]

HEADERS_OA = [
    "Date of Birth",
    "Age completed",
    *oldage_slab_excel_headers(),
    "Payable Basic (Higher of M-I & M-II)",
    "Payable Basic + Old Age Benefit",
    "Current rate kind",
    "Current rate %",
    "Current rate valid from",
    "Current rate valid to",
]


def _report_scale(row) -> str:
    from methodology2.services.class1_2_service import pay_scale_for_report

    return pay_scale_for_report(
        getattr(row, "scale", None),
        getattr(row, "retirement_date", None),
    )


def _num(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _fmt_date(value):
    if value is None or value == "":
        return ""
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    text = str(value).strip()
    if len(text) >= 10 and text[4] == "-":
        try:
            return date.fromisoformat(text[:10]).strftime("%d-%m-%Y")
        except ValueError:
            return text
    return text


def _class_label(category):
    cat = str(category or "").strip()
    mapping = {"1": "I", "2": "II", "3": "III", "4": "IV"}
    return mapping.get(cat, cat)


def _display_name(row) -> str:
    if getattr(row, "is_employee_pension", False):
        return (row.name or row.wage_emp_name or "").strip()
    return (row.pensioner_name or row.name or "").strip()


def _beneficial(m1, m2):
    """Label which methodology is higher; amount = M2 − M1 (negative if M-I higher), as PDF."""
    if m1 is None and m2 is None:
        return "", None
    if m1 is None:
        return "M-II", None
    if m2 is None:
        return "M-I", None
    try:
        a = float(m1)
        b = float(m2)
    except (TypeError, ValueError):
        return "", None
    diff = round(b - a, 2)
    if b > a:
        return "M-II", diff
    if a > b:
        return "M-I", diff  # negative
    return "Equal", 0.0


def _higher_of(m1, m2):
    if m1 is None and m2 is None:
        return None
    if m1 is None:
        return m2
    if m2 is None:
        return m1
    return max(m1, m2)


def _comparison(row) -> dict:
    cmp_ = getattr(row, "pension_comparison", None) or {}
    return cmp_ if isinstance(cmp_, dict) else {}


def _columns(row) -> list[dict]:
    cols = _comparison(row).get("columns") or []
    return [c for c in cols if isinstance(c, dict)]


def _m2_m1_for_cpi(row, cpi: int):
    """Prefer last matching print column for CPI (handles enhanced → family).

    Returns (m1, m2, signed_diff) with signed_diff = M2 − M1 (PDF / may be negative).
    """
    cols = [c for c in _columns(row) if int(c.get("cpi") or 0) == cpi]
    if cols:
        # Prefer current / family / last
        today = date.today().isoformat()
        current = None
        for c in cols:
            pf = str(c.get("period_from") or "")[:10]
            pt = str(c.get("period_to") or "")[:10]
            if pf and pf <= today and (not pt or pt >= today):
                current = c
                break
        col = current or cols[-1]
        m1 = _num(col.get("m1"))
        m2 = _num(
            col.get("m2_base")
            if col.get("m2_base") is not None
            else col.get("m2")
        )
        if m1 is not None and m2 is not None:
            diff = round(m2 - m1, 2)
        else:
            diff = _num(col.get("diff"))
        return m1, m2, diff

    m1 = _num(getattr(row, "m1_family_pension_277" if cpi == 277 else "m1_family_pension_359", None))
    if getattr(row, "is_employee_pension", False):
        m2 = _num(getattr(row, "m2_pension_277" if cpi == 277 else "m2_pension_359", None))
    else:
        m2 = _num(
            getattr(
                row,
                "m2_family_pension_277" if cpi == 277 else "m2_family_pension_359",
                None,
            )
        )
    diff = round(m2 - m1, 2) if m1 is not None and m2 is not None else None
    return m1, m2, diff


def _oldage_dict(row) -> dict:
    oa = getattr(row, "oldage_benefit", None) or {}
    return oa if isinstance(oa, dict) else {}


def _oa_from_columns(row) -> dict:
    """Pick current OA column amounts from pension_comparison when present."""
    today = date.today().isoformat()
    best = None
    for c in _columns(row):
        if c.get("oa_benefit") is None and c.get("oldage_percent") in (None, 0, 0.0):
            continue
        pf = str(c.get("period_from") or "")[:10]
        pt = str(c.get("period_to") or "")[:10]
        if pf and pf <= today and (not pt or pt >= today):
            best = c
            break
        if best is None:
            best = c
    if not best:
        return {}
    return {
        "oa_percent": _num(best.get("oldage_percent")),
        "oa_benefit_m1": _num(best.get("oa_benefit_m1")),
        "oa_benefit_m2": _num(best.get("oa_benefit_m2")),
        "oa_benefit": _num(best.get("oa_benefit")),
        "payable_basic": _num(best.get("payable_basic")),
        "payable_with_oa": _num(
            best.get("revised_total_basic") or best.get("payable_with_oa")
        ),
        "rate_kind": best.get("kind"),
        "valid_from": best.get("period_from"),
        "valid_to": best.get("period_to"),
    }


def _consol_rate_current(emp_cd: str, source_app: str) -> M2ConsolRate | None:
    case = (
        M2ConsolCase.objects.filter(emp_cd=emp_cd, source_app=source_app)
        .prefetch_related("rates")
        .first()
    )
    if not case:
        return None
    rates = list(case.rates.all())
    for r in rates:
        if r.is_current:
            return r
    today = date.today()
    for r in rates:
        if r.valid_from <= today and (r.valid_to is None or r.valid_to >= today):
            return r
    return rates[-1] if rates else None


def _oa_row_values(row, *, source_app: str) -> list:
    oa = _oldage_dict(row)
    from_cols = _oa_from_columns(row)
    rate = _consol_rate_current(str(row.emp_cd), source_app)

    dob = getattr(row, "date_of_birth", None) or oa.get("dob") or _comparison(row).get(
        "date_of_birth"
    )
    age = oa.get("age_completed")

    oa_m1 = from_cols.get("oa_benefit_m1")
    oa_m2 = from_cols.get("oa_benefit_m2")
    oa_ben = _num(oa.get("additional_pension")) or from_cols.get("oa_benefit")
    if oa_ben is None and rate is not None:
        oa_ben = _num(rate.oa_benefit)
    if oa_ben is None and oa_m1 is not None and oa_m2 is not None:
        oa_ben = max(oa_m1, oa_m2)
    elif oa_ben is None:
        oa_ben = oa_m2 if oa_m2 is not None else oa_m1

    m1_359, m2_359, _ = _m2_m1_for_cpi(row, 359)
    payable = from_cols.get("payable_basic")
    if payable is None and rate is not None:
        payable = _num(rate.payable_amount)
    if payable is None:
        payable = _higher_of(m1_359, m2_359)
        if payable is None:
            m1_277, m2_277, _ = _m2_m1_for_cpi(row, 277)
            payable = _higher_of(m1_277, m2_277)

    with_oa = from_cols.get("payable_with_oa")
    if with_oa is None and rate is not None:
        with_oa = _num(rate.payable_with_oa)
    if with_oa is None and payable is not None and oa_ben is not None:
        with_oa = round(float(payable) + float(oa_ben), 2)
    elif with_oa is None:
        with_oa = payable

    kind = ""
    rate_pct = None
    vf = ""
    vt = ""
    if rate is not None:
        kind = rate.rate_kind or ""
        rate_pct = _num(rate.rate_pct)
        vf = _fmt_date(rate.valid_from)
        vt = _fmt_date(rate.valid_to) or "Open"
    elif from_cols.get("rate_kind"):
        kind = str(from_cols.get("rate_kind") or "").upper()
        vf = _fmt_date(from_cols.get("valid_from"))
        vt = _fmt_date(from_cols.get("valid_to")) or ("Open" if from_cols.get("valid_from") else "")

    return [
        _fmt_date(dob),
        age if age is not None else "",
        *oldage_slab_excel_values(_columns(row)),
        payable,
        with_oa,
        kind,
        rate_pct,
        vf,
        vt,
    ]


def _enhanced_row_values(row, *, source_app: str) -> list:
    cmp_ = _comparison(row)
    case = M2ConsolCase.objects.filter(
        emp_cd=str(row.emp_cd), source_app=source_app
    ).first()
    enhanced_upto = None
    family_from = None
    if case:
        enhanced_upto = case.enhanced_upto
        family_from = case.family_from
    if not enhanced_upto:
        enhanced_upto = cmp_.get("double_fpension_upto") or cmp_.get("rate_cutover")
    if not family_from and enhanced_upto:
        try:
            d = (
                enhanced_upto
                if isinstance(enhanced_upto, date)
                else date.fromisoformat(str(enhanced_upto)[:10])
            )
            family_from = d + timedelta(days=1)
        except ValueError:
            family_from = None
    has_split = bool(
        cmp_.get("has_death_split")
        or (case and case.case_type in (M2ConsolCase.CASE_DIH, M2ConsolCase.CASE_DIED_AFTER_RET))
    )
    return [
        _fmt_date(enhanced_upto),
        _fmt_date(family_from),
        "Yes" if has_split else "No",
    ]


def _row_values(
    sl_no: int,
    row,
    *,
    include_oldage: bool,
    include_enhanced: bool,
    source_app: str,
):
    m1_277, m2_277, amt_277 = _m2_m1_for_cpi(row, 277)
    m1_359, m2_359, amt_359 = _m2_m1_for_cpi(row, 359)
    ben_277, _ = _beneficial(m1_277, m2_277)
    ben_359, _ = _beneficial(m1_359, m2_359)
    ret = row.retirement_date
    dod = getattr(row, "date_of_death", None) or _comparison(row).get("date_of_death")
    values = [
        sl_no,
        _display_name(row),
        row.case_no or "",
        row.roll_no or "",
        row.emp_cd,
        ret.strftime("%d-%m-%Y") if ret else "",
        _class_label(row.category),
        row.designation or "",
        _num(row.last_pay),
        _report_scale(row),
        _fmt_date(dod),
        m1_277,
        m2_277,
        ben_277,
        amt_277,
        m1_359,
        m2_359,
        ben_359,
        amt_359,
    ]
    if include_enhanced:
        values.extend(_enhanced_row_values(row, source_app=source_app))
    if include_oldage:
        values.extend(_oa_row_values(row, source_app=source_app))
    return values


def report_filename(
    from_date: date | None = None,
    to_date: date | None = None,
    *,
    include_oldage: bool = False,
) -> str:
    prefix = "M2_OldAge_consolidation_report" if include_oldage else "M2_consolidation_report"
    if from_date and to_date:
        return f"{prefix}_{from_date:%Y%m%d}_{to_date:%Y%m%d}.xlsx"
    return f"{prefix}_{date.today():%Y-%m-%d}.xlsx"


def _inclusive_updated_range(from_date: date, to_date: date):
    """
    Day bounds for updated_at filtering.

    Avoid __date / CONVERT_TZ: MySQL often lacks timezone tables, so
    DATE(CONVERT_TZ(...)) becomes NULL and returns zero rows.
    """
    start = timezone.make_aware(datetime.combine(from_date, time.min))
    end_excl = timezone.make_aware(
        datetime.combine(to_date + timedelta(days=1), time.min)
    )
    return start, end_excl


def build_consolidation_report_bytes(
    from_date: date,
    to_date: date,
    *,
    include_oldage: bool = False,
    include_enhanced: bool = True,
    model=None,
    source_app: str | None = None,
) -> tuple[bytes, int]:
    """
    Excel export filtered by consolidation.updated_at date range
    (inclusive calendar days in Django TIME_ZONE).

    include_oldage: add OA % / benefit / payable+OA columns (OA app report).
    include_enhanced: add enhanced-upto / family-from columns.
    """
    if from_date > to_date:
        raise ValueError("from_date cannot be after to_date")

    Model = model or Methodology2Consolidation
    app = source_app or (
        M2ConsolCase.SOURCE_OA if include_oldage else M2ConsolCase.SOURCE_MET2
    )

    start, end_excl = _inclusive_updated_range(from_date, to_date)
    rows = list(
        Model.objects.filter(
            updated_at__gte=start,
            updated_at__lt=end_excl,
        ).order_by("case_no", "emp_cd")
    )

    headers = list(HEADERS_BASE)
    if include_enhanced:
        headers.extend(HEADERS_ENHANCED)
    if include_oldage:
        headers.extend(HEADERS_OA)

    wb = Workbook()
    ws = wb.active
    ws.title = "OA Consolidation" if include_oldage else "M1 vs M2 Report"
    ws.append(headers)

    header_font = Font(bold=True)
    header_align = Alignment(
        wrap_text=True, vertical="center", horizontal="center"
    )
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = header_align

    for idx, row in enumerate(rows, start=1):
        ws.append(
            _row_values(
                idx,
                row,
                include_oldage=include_oldage,
                include_enhanced=include_enhanced,
                source_app=app,
            )
        )

    for col_idx, header in enumerate(headers, start=1):
        width = min(max(len(header) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), len(rows)
